#!/usr/bin/env python3
"""Prepare dropped documents for knowledge extraction and check the result.

Commands (JSON on stdout, exit 1 on failure):

  scan <folder>        hash, secret-scan, extract text, and chunk every document
  chunks <doc-id>      print a document's chunks for reading
  record <doc-id>      record curated, merged, or rejected for one document
  check                every document resolved and every written concept well-formed

State lives in <workspace>/.knowledge-extract/: manifest.json plus derived/<sha256>.md
(extracted text) and derived/<sha256>.chunks.json. Unchanged files are skipped on a re-scan;
a changed file reopens and lists the concepts it fed. Standard library plus PyYAML; .pdf
needs pypdf and .xlsx needs openpyxl, and a missing parser skips the file with a reason.
"""

from __future__ import annotations

import argparse
import hashlib
import html.parser
import json
import os
import re
import sys
import tempfile
import zipfile
from collections import deque
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from urllib.parse import unquote
from xml.etree import ElementTree

try:
    import yaml
except ImportError as exc:  # pragma: no cover
    raise SystemExit("PyYAML is required: python -m pip install pyyaml") from exc

STATE_DIR = ".knowledge-extract"
MAX_FILE_BYTES = 50_000_000
MAX_ARCHIVE_BYTES = 200_000_000
MAX_CHARS = 2_000_000
MAX_ROWS = 2_000
MAX_CHUNK_CHARS = 1_800
MIN_SECTION_CHARS = 200
ROWS_PER_CHUNK = 40
LINES_PER_CHUNK = 80

DOCUMENT = {".md", ".markdown", ".txt", ".rst"}
TABULAR = {".csv", ".tsv"}
CODE = {".json", ".yaml", ".yml", ".sql", ".xml", ".ini", ".toml", ".log", ".py", ".sh", ".ps1"}
OFFICE = {".docx", ".pptx", ".xlsx", ".xlsm", ".pdf", ".html", ".htm"}
SUPPORTED = DOCUMENT | TABULAR | CODE | OFFICE
IGNORED_DIRS = {".git", STATE_DIR, "__pycache__", "node_modules", ".venv"}

OPEN = "pending"
RESOLVED = {"curated", "merged", "rejected"}
LIMITS = {"refused", "unsupported", "skipped", "failed", "empty"}

# Durability: a first guess at whether a chunk holds reusable knowledge. "unsure" needs the
# model's judgement. The first matching rule wins.
_EPHEMERAL_SECTION = re.compile(
    r"\b(agenda|attendees|participants|housekeeping|logistics|status update|next steps|action items)\b",
    re.IGNORECASE,
)
_UNSURE_NAME = re.compile(
    r"transcript|meeting|minutes|notes|session|stand-?up|retro|incident|rca|post-?mortem|"
    r"e-?mail|thread|chat|slack|teams|interview|status|update|report",
    re.IGNORECASE,
)
_DURABLE_NAME = re.compile(
    r"runbook|sop|playbook|procedure|guideline|policy|standard|checklist|escalation|on-?call|"
    r"owner|raci|contact|glossary|definition|dictionary|catalog|schema|architecture|design|"
    r"sla|slo|kpi|template|how-?to|guide|faq|onboarding|handbook|manual|reference|rules?",
    re.IGNORECASE,
)

SECRET_PATTERNS: tuple[tuple[str, re.Pattern[str]], ...] = (
    ("private_key", re.compile(r"-----BEGIN (?:RSA |EC |OPENSSH |DSA )?PRIVATE KEY-----")),
    ("aws_access_key", re.compile(r"\b(?:AKIA|ASIA)[A-Z0-9]{16}\b")),
    ("databricks_pat", re.compile(r"\bdapi[a-fA-F0-9]{24,}\b")),
    ("openai_key", re.compile(r"\bsk-(?:proj-|ant-)?[A-Za-z0-9_-]{20,}\b")),
    ("github_token", re.compile(r"\bgh[pousr]_[A-Za-z0-9]{30,}\b")),
    ("slack_token", re.compile(r"\bxox[abposr]-[A-Za-z0-9-]{10,}\b")),
    ("bearer_token", re.compile(r"(?i)\bbearer\s+[A-Za-z0-9_\-.=]{30,}")),
    (
        "assigned_secret",
        re.compile(
            r"(?i)(?:^|[\s,{])(?:[a-z0-9]+[_-])?"
            r"(?:password|passwd|pwd|api[_-]?key|access[_-]?token|client[_-]?secret|secret[_-]?key)"
            r"\s*[:=]\s*[\"']?(?!\$\{|\{\{|<|REDACTED|CHANGEME|replace-me|\*{3})"
            r"[A-Za-z0-9_./+=:@!#$%^&-]{12,}"
        ),
    ),
)

FRONTMATTER = re.compile(r"\A---\r?\n(.*?)\r?\n---(?:\r?\n|\Z)", re.DOTALL)
HEADING = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")
LINK = re.compile(r"(?<!!)\[[^\]]*\]\(([^)\s]+)(?:\s+\"[^\"]*\")?\)")
FOOTNOTE_REF = re.compile(r"\[\^([^\]\s]+)\](?!:)")
FOOTNOTE_DEF = re.compile(r"^\[\^([^\]\s]+)\]:", re.MULTILINE)
DOC_ID = re.compile(r"^doc-[0-9a-f]{12}$")
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


class ExtractError(Exception):
    """A document could not be read; the message is safe to show."""


class MissingParser(ExtractError):
    """An optional parser is not installed."""


# --------------------------------------------------------------------------- helpers


def now() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def write_atomic(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    # A short temporary name: derived files are named by their 64-character digest already.
    handle, temp = tempfile.mkstemp(dir=path.parent, prefix=".tmp-", suffix=".part")
    try:
        with os.fdopen(handle, "w", encoding="utf-8", newline="\n") as stream:
            stream.write(text)
        os.replace(temp, path)
    except BaseException:
        Path(temp).unlink(missing_ok=True)
        raise


def within(path: Path, parent: Path) -> bool:
    try:
        path.resolve().relative_to(parent.resolve())
        return True
    except ValueError:
        return False


def scan_secrets(text: str) -> list[dict[str, Any]]:
    findings = []
    for number, line in enumerate(text.splitlines(), start=1):
        for name, pattern in SECRET_PATTERNS:
            if pattern.search(line):
                findings.append({"line": number, "detector": name})
    return findings


def knowledge_root(workspace: Path, override: str | None) -> Path:
    value = override
    if value is None:
        config = workspace / "operations.yaml"
        if config.is_file():
            data = yaml.safe_load(config.read_text(encoding="utf-8")) or {}
            knowledge = data.get("knowledge") if isinstance(data, dict) else None
            if isinstance(knowledge, dict) and isinstance(knowledge.get("path"), str):
                value = knowledge["path"]
    root = (workspace / (value or "knowledge")).resolve()
    if not within(root, workspace):
        raise SystemExit(f"knowledge path must stay inside the workspace: {value}")
    return root


class Manifest:
    def __init__(self, workspace: Path) -> None:
        self.workspace = workspace
        self.root = workspace / STATE_DIR
        self.path = self.root / "manifest.json"
        if self.path.is_file():
            self.data = json.loads(self.path.read_text(encoding="utf-8"))
        else:
            self.data = {"version": 1, "documents": {}}

    @property
    def documents(self) -> dict[str, dict[str, Any]]:
        return self.data["documents"]

    def get(self, doc_id: str) -> dict[str, Any]:
        if not DOC_ID.fullmatch(doc_id) or doc_id not in self.documents:
            raise SystemExit(f"unknown document id: {doc_id}")
        return self.documents[doc_id]

    def derived(self, digest: str, suffix: str) -> Path:
        return self.root / "derived" / f"{digest}{suffix}"

    def save(self) -> None:
        write_atomic(self.path, json.dumps(self.data, indent=2, sort_keys=True) + "\n")


# --------------------------------------------------------------------------- extraction


def decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ExtractError("text encoding is not supported")


def _open_archive(path: Path) -> zipfile.ZipFile:
    try:
        archive = zipfile.ZipFile(path)
    except zipfile.BadZipFile as exc:
        raise ExtractError("not a valid Office file") from exc
    if sum(item.file_size for item in archive.infolist()) > MAX_ARCHIVE_BYTES:
        archive.close()
        raise ExtractError(f"archive expands beyond {MAX_ARCHIVE_BYTES} bytes")
    return archive


def _xml(data: bytes) -> ElementTree.Element:
    try:
        return ElementTree.fromstring(data)
    except ElementTree.ParseError as exc:
        raise ExtractError(f"unreadable document XML: {exc}") from exc


def _natural(value: str) -> list[Any]:
    return [int(part) if part.isdigit() else part for part in re.split(r"(\d+)", value)]


def extract_docx(path: Path) -> str:
    with _open_archive(path) as archive:
        if "word/document.xml" not in archive.namelist():
            raise ExtractError("Word file has no document body")
        root = _xml(archive.read("word/document.xml"))
        lines: list[str] = []
        for paragraph in root.iter(f"{W}p"):
            text = "".join(node.text or "" for node in paragraph.iter(f"{W}t")).strip()
            if not text:
                continue
            style = paragraph.find(f"{W}pPr/{W}pStyle")
            name = (style.get(f"{W}val") or "") if style is not None else ""
            level = re.fullmatch(r"(?i)heading\s?(\d)", name)
            if level:
                lines.append(f"\n{'#' * min(int(level.group(1)), 6)} {text}")
            elif name.lower() == "title":
                lines.append(f"\n# {text}")
            elif paragraph.find(f"{W}pPr/{W}numPr") is not None or name.lower().startswith("list"):
                lines.append(f"- {text}")
            else:
                lines.append(text)
        if "word/comments.xml" in archive.namelist():
            comments = [
                "".join(node.text or "" for node in comment.iter(f"{W}t")).strip()
                for comment in _xml(archive.read("word/comments.xml")).iter(f"{W}comment")
            ]
            if any(comments):
                lines.append("\n## Comments")
                lines.extend(f"- {comment}" for comment in comments if comment)
    return "\n".join(lines).strip()


def extract_pptx(path: Path) -> str:
    def paragraphs(data: bytes) -> list[str]:
        out = []
        for paragraph in _xml(data).iter(f"{A}p"):
            text = "".join(node.text or "" for node in paragraph.iter(f"{A}t")).strip()
            if text:
                out.append(text)
        return out

    with _open_archive(path) as archive:
        names = archive.namelist()
        slides = sorted((n for n in names if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)), key=_natural)
        parts: list[str] = []
        for number, name in enumerate(slides, start=1):
            texts = paragraphs(archive.read(name))
            title = texts[0] if texts else "(no text)"
            parts.append(f"\n## Slide {number}: {title}")
            parts.extend(texts[1:])
            slide_number = re.search(r"(\d+)\.xml$", name).group(1)
            notes = f"ppt/notesSlides/notesSlide{slide_number}.xml"
            if notes in names:
                note_text = [t for t in paragraphs(archive.read(notes)) if not t.isdigit()]
                if note_text:
                    parts.append("\n### Speaker notes")
                    parts.extend(note_text)
    return "\n".join(parts).strip()


def extract_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise MissingParser("PDF needs pypdf: python -m pip install pypdf") from exc
    try:
        reader = PdfReader(str(path))
        pages = [page.extract_text() or "" for page in reader.pages]
    except Exception as exc:  # pypdf raises many types for damaged files
        raise ExtractError(f"unreadable PDF: {type(exc).__name__}") from exc
    return "\n\n".join(f"## Page {n}\n{text.strip()}" for n, text in enumerate(pages, start=1) if text.strip())


def extract_xlsx(path: Path) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise MissingParser("Excel needs openpyxl: python -m pip install openpyxl") from exc
    _open_archive(path).close()
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise ExtractError(f"unreadable workbook: {type(exc).__name__}") from exc
    parts: list[str] = []
    remaining = MAX_ROWS
    try:
        for sheet in workbook.worksheets:
            parts.append(f"\n## Sheet: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                if remaining <= 0:
                    parts.append("[row limit reached]")
                    return "\n".join(parts).strip()
                cells = ["" if value is None else str(value) for value in row[:100]]
                if any(cells):
                    parts.append(" | ".join(cells))
                    remaining -= 1
    finally:
        workbook.close()
    return "\n".join(parts).strip()


class _HtmlText(html.parser.HTMLParser):
    BLOCK = {"p", "div", "br", "tr", "section", "article", "table", "ul", "ol", "pre", "blockquote"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.skip = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in {"script", "style", "noscript", "head"}:
            self.skip += 1
        elif re.fullmatch(r"h[1-6]", tag):
            self.parts.append(f"\n\n{'#' * int(tag[1])} ")
        elif tag == "li":
            self.parts.append("\n- ")
        elif tag in {"td", "th"}:
            self.parts.append(" | ")
        elif tag in self.BLOCK:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style", "noscript", "head"}:
            self.skip = max(0, self.skip - 1)
        elif re.fullmatch(r"h[1-6]", tag) or tag in self.BLOCK:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self.skip:
            self.parts.append(re.sub(r"\s+", " ", data))


def extract_html(path: Path) -> str:
    parser = _HtmlText()
    parser.feed(decode(path.read_bytes()))
    lines = [line.strip() for line in "".join(parser.parts).splitlines()]
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in DOCUMENT | TABULAR | CODE:
        return decode(path.read_bytes())
    if suffix == ".docx":
        return extract_docx(path)
    if suffix == ".pptx":
        return extract_pptx(path)
    if suffix == ".pdf":
        return extract_pdf(path)
    if suffix in {".xlsx", ".xlsm"}:
        return extract_xlsx(path)
    if suffix in {".html", ".htm"}:
        return extract_html(path)
    raise ExtractError(f"unsupported format: {suffix or '[no extension]'}")


# --------------------------------------------------------------------------- chunking


def kind_of(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in TABULAR:
        return "tabular"
    if suffix in CODE:
        return "code"
    return "document"


def durability(name: str, section: str, kind: str, *, header: bool = False) -> tuple[str, str]:
    stem = Path(name).stem
    if kind == "tabular":
        return ("yes", "column header describes structure") if header else ("no", "row values")
    if _EPHEMERAL_SECTION.search(section):
        return "no", "meeting logistics or status section"
    if _UNSURE_NAME.search(stem):
        return "unsure", "notes, discussion, or incident narrative; judge what is reusable"
    if _DURABLE_NAME.search(stem) or _DURABLE_NAME.search(section):
        return "yes", "procedure, definition, ownership, or reference material"
    if kind == "code":
        return "yes", "configuration or code"
    return "unsure", "no rule matched"


def _split_long(text: str) -> list[str]:
    if len(text) <= MAX_CHUNK_CHARS:
        return [text]
    pieces: list[str] = []
    current: list[str] = []
    size = 0
    for paragraph in re.split(r"\n\s*\n", text):
        for piece in [paragraph] if len(paragraph) <= MAX_CHUNK_CHARS else paragraph.splitlines():
            if size + len(piece) > MAX_CHUNK_CHARS and current:
                pieces.append("\n\n".join(current))
                current, size = [], 0
            current.append(piece[:MAX_CHUNK_CHARS])
            size += len(piece) + 2
    if current:
        pieces.append("\n\n".join(current))
    return pieces


def _sections(content: str) -> list[tuple[str, str]]:
    trail: list[tuple[int, str]] = []
    body: list[str] = []
    out: list[tuple[str, str]] = []

    def emit() -> None:
        text = "\n".join(body).strip()
        if text:
            out.append((" > ".join(title for _, title in trail), text))

    for line in content.splitlines():
        match = HEADING.match(line)
        if match:
            emit()
            body = []
            level = len(match.group(1))
            trail = [(depth, title) for depth, title in trail if depth < level] + [(level, match.group(2))]
        else:
            body.append(line)
    emit()
    return out


def chunk(name: str, content: str, kind: str) -> list[dict[str, Any]]:
    chunks: list[dict[str, Any]] = []

    def add(section: str, text: str, *, header: bool = False) -> None:
        durable, reason = durability(name, section, kind, header=header)
        chunks.append({"n": len(chunks), "section": section, "durable": durable, "reason": reason, "text": text})

    if kind == "tabular":
        lines = [line for line in content.splitlines() if line.strip()]
        if lines:
            header, rows = lines[0], lines[1:]
            add("columns", f"Columns: {header}\nRows: {len(rows)}\n" + "\n".join(rows[:5]), header=True)
            for start in range(0, len(rows), ROWS_PER_CHUNK):
                add(f"rows {start + 1}-{min(start + ROWS_PER_CHUNK, len(rows))}",
                    header + "\n" + "\n".join(rows[start:start + ROWS_PER_CHUNK]))
        return chunks
    if kind == "code":
        lines = content.splitlines()
        for start in range(0, len(lines), LINES_PER_CHUNK):
            block = "\n".join(lines[start:start + LINES_PER_CHUNK]).strip()
            if block:
                add(f"lines {start + 1}-{min(start + LINES_PER_CHUNK, len(lines))}", block[: MAX_CHUNK_CHARS * 2])
        return chunks

    pending_section, pending_text = "", ""
    for section, body in _sections(content):
        if pending_text and len(pending_text) < MIN_SECTION_CHARS:
            # Fold a fragment into the next section and keep its heading inline.
            prefix = f"{pending_section.rsplit(' > ', 1)[-1]}\n" if pending_section else ""
            body = f"{prefix}{pending_text}\n\n{body}"
        elif pending_text:
            for piece in _split_long(pending_text):
                add(pending_section, piece)
        pending_section, pending_text = section, body
    if pending_text:
        for piece in _split_long(pending_text):
            add(pending_section, piece)
    return chunks


def document_durability(chunks: list[dict[str, Any]]) -> str:
    values = {item["durable"] for item in chunks}
    return "yes" if "yes" in values else "unsure" if "unsure" in values else "no"


# --------------------------------------------------------------------------- commands


def _discover(folder: Path) -> list[Path]:
    files = []
    for path in sorted(folder.rglob("*")):
        relative = path.relative_to(folder)
        if any(part in IGNORED_DIRS or part.startswith(".") for part in relative.parts):
            continue
        if path.is_file() and within(path, folder):
            files.append(path)
    return files


def scan(folder: Path, workspace: Path) -> dict[str, Any]:
    folder = folder.resolve()
    if not folder.is_dir():
        raise SystemExit(f"folder not found: {folder}")
    if not within(folder, workspace):
        raise SystemExit("the document folder must be inside the workspace")
    manifest = Manifest(workspace)
    report: dict[str, Any] = {"folder": folder.relative_to(workspace).as_posix(), "new": [], "changed": [],
                              "unchanged": 0, "removed": [], "limits": [], "pending": []}
    seen: set[str] = set()

    for path in _discover(folder):
        relative = path.relative_to(workspace).as_posix()
        doc_id = "doc-" + hashlib.sha256(relative.encode("utf-8")).hexdigest()[:12]
        seen.add(doc_id)
        previous = manifest.documents.get(doc_id)
        size = path.stat().st_size
        digest = sha256_file(path)
        if previous and previous["sha256"] == digest and previous["status"] not in {"skipped", "failed"}:
            report["unchanged"] += 1
            continue

        record: dict[str, Any] = {"path": relative, "sha256": digest, "bytes": size, "kind": kind_of(path),
                                  "scanned_at": now(), "status": OPEN, "concepts": [], "reason": ""}
        if previous and previous.get("concepts"):
            record["revisit"] = sorted(set(previous["concepts"]) | set(previous.get("revisit", [])))
        try:
            if path.suffix.lower() not in SUPPORTED:
                raise ExtractError(f"unsupported format: {path.suffix or '[no extension]'}")
            if size > MAX_FILE_BYTES:
                raise MissingParser(f"larger than {MAX_FILE_BYTES} bytes")
            text = extract_text(path)
        except MissingParser as exc:
            record.update(status="skipped", reason=str(exc))
        except ExtractError as exc:
            status = "unsupported" if str(exc).startswith("unsupported") else "failed"
            record.update(status=status, reason=str(exc))
        else:
            findings = scan_secrets(text)
            if findings:
                where = ", ".join(f"line {f['line']} ({f['detector']})" for f in findings[:10])
                record.update(status="refused", reason=f"possible secret at {where}; redact and upload again")
            elif not text.strip():
                record.update(status="empty", reason="no extractable text (scanned PDF or image?)")
            else:
                if len(text) > MAX_CHARS:
                    text = text[:MAX_CHARS] + "\n\n[character limit reached]\n"
                chunks = chunk(path.name, text, record["kind"])
                record["durable"] = document_durability(chunks)
                record["chunks"] = [{k: v for k, v in item.items() if k != "text"} | {"chars": len(item["text"])}
                                    for item in chunks]
                write_atomic(manifest.derived(digest, ".md"),
                             f"# {path.name}\n\n- Resource: `evidence://sha256/{digest}/{path.name}`\n"
                             f"- Source: `{relative}`\n\n## Extracted content\n\n{text}\n")
                write_atomic(manifest.derived(digest, ".chunks.json"), json.dumps(chunks, indent=1) + "\n")

        manifest.documents[doc_id] = record
        entry = {"id": doc_id, "path": relative}
        if previous:
            report["changed"].append(entry | {"revisit": record.get("revisit", [])})
        else:
            report["new"].append(entry)

    for doc_id, record in manifest.documents.items():
        in_folder = within(workspace / record["path"], folder)
        if in_folder and doc_id not in seen and record["status"] != "removed":
            if record["status"] in RESOLVED and record["concepts"]:
                continue  # keep what was learned; the knowledge stays reviewable on its own
            record.update(status="removed", reason="file no longer in the folder")
            report["removed"].append({"id": doc_id, "path": record["path"]})

    manifest.save()
    order = {"yes": 0, "unsure": 1, "no": 2}
    for doc_id, record in sorted(manifest.documents.items(),
                                 key=lambda kv: (order.get(kv[1].get("durable", "no"), 3), kv[1]["path"])):
        if record["status"] == OPEN:
            report["pending"].append({"id": doc_id, "path": record["path"], "kind": record["kind"],
                                      "durable": record.get("durable"), "chunks": len(record.get("chunks", [])),
                                      "revisit": record.get("revisit", [])})
        elif record["status"] in LIMITS:
            report["limits"].append({"id": doc_id, "path": record["path"], "status": record["status"],
                                     "reason": record["reason"]})
    return report


def chunks_command(workspace: Path, doc_id: str, number: int | None) -> dict[str, Any]:
    manifest = Manifest(workspace)
    record = manifest.get(doc_id)
    source = manifest.derived(record["sha256"], ".chunks.json")
    if not source.is_file():
        raise SystemExit(f"{doc_id} has no extracted text ({record['status']}: {record['reason']})")
    items = json.loads(source.read_text(encoding="utf-8"))
    if number is not None:
        if not 0 <= number < len(items):
            raise SystemExit(f"{doc_id} has chunks 0-{len(items) - 1}")
        items = [items[number]]
    return {"id": doc_id, "path": record["path"], "sha256": record["sha256"], "status": record["status"],
            "chunks": items}


def record_command(workspace: Path, knowledge: Path, doc_id: str, result: str,
                   concepts: list[str], reason: str | None) -> dict[str, Any]:
    manifest = Manifest(workspace)
    record = manifest.get(doc_id)
    if record["status"] not in {OPEN} | RESOLVED:
        raise SystemExit(f"{doc_id} is {record['status']} and cannot be recorded ({record['reason']})")
    cleaned = []
    for value in concepts:
        relative = value.replace("\\", "/").lstrip("/")
        target = (knowledge / relative).resolve()
        if not within(target, knowledge) or target.suffix != ".md":
            raise SystemExit(f"concept must be a .md path inside the knowledge folder: {value}")
        if not target.is_file():
            raise SystemExit(f"concept does not exist yet: {relative}")
        cleaned.append(target.relative_to(knowledge).as_posix())
    if result in {"curated", "merged"} and not cleaned:
        raise SystemExit(f"--result {result} needs at least one --concept")
    if result == "rejected" and not (reason and reason.strip()):
        raise SystemExit("--result rejected needs a --reason")
    record.update(status=result, concepts=sorted(set(cleaned)), reason=(reason or "").strip(), recorded_at=now())
    record.pop("revisit", None)
    manifest.save()
    return {"id": doc_id, "path": record["path"], "status": result, "concepts": record["concepts"]}


def _frontmatter(text: str) -> tuple[dict[str, Any] | None, str, str | None]:
    match = FRONTMATTER.match(text)
    if not match:
        return None, text, "missing YAML frontmatter"
    try:
        data = yaml.safe_load(match.group(1))
    except yaml.YAMLError as exc:
        return None, text, f"invalid YAML frontmatter: {str(exc).splitlines()[0]}"
    if not isinstance(data, dict):
        return None, text, "frontmatter must be a mapping"
    return data, text[match.end():], None


def _link_targets(path: Path, body: str, knowledge: Path) -> list[tuple[str, Path]]:
    targets = []
    for raw in LINK.findall(body):
        clean = unquote(raw.split("#", 1)[0].strip())
        if not clean or re.match(r"^[a-z][a-z0-9+.-]*:", clean, re.IGNORECASE):
            continue
        base = knowledge / clean.lstrip("/") if clean.startswith("/") else path.parent / clean
        targets.append((raw, base.resolve()))
    return targets


def _reachable(knowledge: Path) -> set[Path]:
    start = knowledge / "index.md"
    if not start.is_file():
        return set()
    seen = {start.resolve()}
    queue = deque([start])
    while queue:
        current = queue.popleft()
        try:
            body = current.read_text(encoding="utf-8")
        except (OSError, UnicodeError):
            continue
        for _, target in _link_targets(current, body, knowledge):
            if target.suffix == ".md" and target.is_file() and within(target, knowledge) and target not in seen:
                seen.add(target)
                queue.append(target)
    return seen


def check(workspace: Path, knowledge: Path) -> dict[str, Any]:
    manifest = Manifest(workspace)
    errors: list[str] = []
    warnings: list[str] = []
    counts: dict[str, int] = {}
    owners: dict[str, list[dict[str, Any]]] = {}
    for doc_id, record in manifest.documents.items():
        counts[record["status"]] = counts.get(record["status"], 0) + 1
        if record["status"] == OPEN:
            errors.append(f"{doc_id} ({record['path']}) is not resolved: record curated, merged, or rejected")
        if record["status"] in {"curated", "merged"}:
            for concept in record["concepts"]:
                owners.setdefault(concept, []).append(record)
        if record["status"] in LIMITS:
            warnings.append(f"{record['path']} was not read: {record['status']} - {record['reason']}")

    reachable = _reachable(knowledge)
    if owners and not (knowledge / "index.md").is_file():
        warnings.append(f"{knowledge.name}/index.md is missing; create the root index and link the new concepts")

    for concept, records in sorted(owners.items()):
        path = knowledge / concept
        label = f"{knowledge.name}/{concept}"
        if not path.is_file():
            errors.append(f"{label} is recorded but does not exist")
            continue
        text = path.read_text(encoding="utf-8")
        data, body, problem = _frontmatter(text)
        if problem:
            errors.append(f"{label}: {problem}")
            continue
        for field in ("type", "title", "description"):
            if not isinstance(data.get(field), str) or not data[field].strip():
                errors.append(f"{label}: frontmatter needs a non-empty {field}")
        status = str(data.get("status", "")).strip().lower()
        if status != "draft" and "<!-- draft" not in body:
            errors.append(f"{label}: new knowledge must be status: draft, or a <!-- draft --> section in an "
                          "existing concept; approval belongs to a human reviewer")
        for record in records:
            name = Path(record["path"]).name
            if record["path"] not in text and name not in text and record["sha256"][:12] not in text:
                errors.append(f"{label}: does not cite its source document {record['path']}")
        for raw, target in _link_targets(path, body, knowledge):
            if not target.exists():
                errors.append(f"{label}: broken link {raw}")
        references = set(FOOTNOTE_REF.findall(body))
        definitions = set(FOOTNOTE_DEF.findall(body))
        for missing in sorted(references - definitions):
            errors.append(f"{label}: footnote [^{missing}] has no definition")
        for finding in scan_secrets(text):
            errors.append(f"{label}: possible secret at line {finding['line']} ({finding['detector']})")
        if path.resolve() not in reachable and path.name != "index.md":
            # A host that links pages when a person keeps them leaves new drafts unlinked on purpose.
            warnings.append(f"{label}: not reachable from the root index.md yet; link it from the right index section")

    return {"ok": not errors, "knowledge": knowledge.relative_to(workspace).as_posix() if within(knowledge, workspace)
            else str(knowledge), "counts": counts, "errors": errors, "warnings": warnings}


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--workspace", type=Path, default=Path.cwd(), help="client workspace root (default: cwd)")
    parser.add_argument("--knowledge", help="knowledge folder relative to the workspace (default: operations.yaml)")
    commands = parser.add_subparsers(dest="command", required=True)
    scan_parser = commands.add_parser("scan", help="read every document in a folder")
    scan_parser.add_argument("folder", type=Path)
    chunks_parser = commands.add_parser("chunks", help="print a document's chunks")
    chunks_parser.add_argument("doc_id")
    chunks_parser.add_argument("--chunk", type=int)
    record_parser = commands.add_parser("record", help="record the outcome for one document")
    record_parser.add_argument("doc_id")
    record_parser.add_argument("--result", required=True, choices=sorted(RESOLVED))
    record_parser.add_argument("--concept", action="append", default=[],
                               help="concept path relative to the knowledge folder; repeatable")
    record_parser.add_argument("--reason")
    commands.add_parser("check", help="verify coverage and the written concepts")
    args = parser.parse_args(argv)

    workspace = args.workspace.resolve()
    if not workspace.is_dir():
        raise SystemExit(f"workspace not found: {workspace}")
    if args.command == "scan":
        folder = args.folder if args.folder.is_absolute() else workspace / args.folder
        result = scan(folder, workspace)
        ok = True
    elif args.command == "chunks":
        result, ok = chunks_command(workspace, args.doc_id, args.chunk), True
    elif args.command == "record":
        knowledge = knowledge_root(workspace, args.knowledge)
        result = record_command(workspace, knowledge, args.doc_id, args.result, args.concept, args.reason)
        ok = True
    else:
        result = check(workspace, knowledge_root(workspace, args.knowledge))
        ok = result["ok"]
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    json.dump(result, sys.stdout, indent=2, ensure_ascii=False)
    sys.stdout.write("\n")
    return 0 if ok else 1


if __name__ == "__main__":
    raise SystemExit(main())
